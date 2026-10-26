"""Export helpers: generate downloadable DOCX bytes from pipeline outputs.

Uses ``python-docx`` to produce a single document that contains the tailored
resume summary and cover letter, ready for ``st.download_button``.
"""

from __future__ import annotations

import io
import datetime


def build_resume_docx(job_title: str, company: str, resume_summary: str, cover_letter: str) -> bytes:
    """Build a DOCX document containing the tailored resume summary and cover letter.

    Args:
        job_title: Job title the candidate is applying for.
        company: Company or employer name.
        resume_summary: AI-generated tailored resume summary (plain text).
        cover_letter: AI-generated cover letter (plain text).

    Returns:
        Raw DOCX bytes suitable for ``st.download_button``.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for DOCX export. Run: pip install python-docx"
        ) from exc

    doc = Document()

    # ── Header ──────────────────────────────────────────────────────────────
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run(f"Application — {job_title} at {company}")
    run.bold = True
    run.font.size = Pt(14)

    date_str = datetime.date.today().strftime("%B %d, %Y")
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(date_str).font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ── Resume Summary ───────────────────────────────────────────────────────
    heading = doc.add_heading("Tailored Resume Summary", level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    if resume_summary:
        doc.add_paragraph(resume_summary)
    else:
        doc.add_paragraph("(No resume summary generated)")

    doc.add_paragraph()  # spacer

    # ── Cover Letter ─────────────────────────────────────────────────────────
    heading2 = doc.add_heading("Cover Letter", level=2)
    for run in heading2.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    if cover_letter:
        for paragraph_text in cover_letter.split("\n"):
            doc.add_paragraph(paragraph_text)
    else:
        doc.add_paragraph("(No cover letter generated)")

    # ── Serialize to bytes ───────────────────────────────────────────────────
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
